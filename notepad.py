import tkinter as tk
from tkinter import ttk, colorchooser, messagebox
from tkinter.scrolledtext import ScrolledText
from tkinter.filedialog import asksaveasfilename, askopenfilename
from docx import Document
from fpdf import FPDF
import datetime
import os

class AdvancedNotepadApp:
    def __init__(self, root):
        self.root = root
        self.root.title("Advanced Notepad")
        self.root.geometry("1000x700")

        self.dark_mode = False
        self.tabs = {}

        # Main Frame and Notebook
        self.main_frame = ttk.Frame(root)
        self.main_frame.pack(fill=tk.BOTH, expand=True)

        self.tab_control = ttk.Notebook(self.main_frame)
        self.tab_control.pack(fill=tk.BOTH, expand=True)
        self.tab_control.bind("<<NotebookTabChanged>>", self.update_status_bar)

        # Menu Bar
        self.menu_bar = tk.Menu(root)
        root.config(menu=self.menu_bar)

        self.file_menu = tk.Menu(self.menu_bar, tearoff=0)
        self.file_menu.add_command(label="New Tab", command=self.new_tab)
        self.file_menu.add_command(label="Open File", command=self.open_file)
        self.file_menu.add_command(label="Save As TXT", command=self.save_as_txt)
        self.file_menu.add_command(label="Save As DOCX", command=self.save_as_docx)
        self.file_menu.add_command(label="Save As PDF", command=self.save_as_pdf)
        self.file_menu.add_separator()
        self.file_menu.add_command(label="Exit", command=root.quit)
        self.menu_bar.add_cascade(label="File", menu=self.file_menu)

        self.edit_menu = tk.Menu(self.menu_bar, tearoff=0)
        self.edit_menu.add_command(label="Undo", command=self.undo)
        self.edit_menu.add_command(label="Redo", command=self.redo)
        self.edit_menu.add_separator()
        self.edit_menu.add_command(label="Cut", command=self.cut)
        self.edit_menu.add_command(label="Copy", command=self.copy)
        self.edit_menu.add_command(label="Paste", command=self.paste)
        self.menu_bar.add_cascade(label="Edit", menu=self.edit_menu)

        self.format_menu = tk.Menu(self.menu_bar, tearoff=0)
        self.format_menu.add_command(label="Bold", command=self.make_bold)
        self.format_menu.add_command(label="Italic", command=self.make_italic)
        self.format_menu.add_command(label="Text Color", command=self.change_text_color)
        self.menu_bar.add_cascade(label="Format", menu=self.format_menu)

        self.view_menu = tk.Menu(self.menu_bar, tearoff=0)
        self.view_menu.add_command(label="Toggle Dark/Light Mode", command=self.toggle_theme)
        self.menu_bar.add_cascade(label="View", menu=self.view_menu)

        # Log Panel
        self.log_panel = tk.Text(root, height=4, bg="#f5f5f5", state=tk.DISABLED)
        self.log_panel.pack(fill=tk.X)

        # Status Bar (Lines, Words, Author)
        status_frame = tk.Frame(root, bg="#e0e0e0")
        status_frame.pack(fill=tk.X)

        self.status_left = tk.Label(status_frame, text="Lines: 0 | Words: 0", anchor=tk.W, bg="#e0e0e0")
        self.status_left.pack(side=tk.LEFT, padx=5)

        self.status_right = tk.Label(status_frame, text="Created by Sayan", anchor=tk.E, bg="#e0e0e0")
        self.status_right.pack(side=tk.RIGHT, padx=5)

        # Add initial tab
        self.new_tab()

    def get_current_tab(self):
        tab_id = self.tab_control.select()
        return self.tabs.get(self.tab_control.nametowidget(tab_id))

    def get_text_widget(self):
        tab = self.get_current_tab()
        return tab['text_widget'] if tab else None

    def log(self, message):
        self.log_panel.config(state=tk.NORMAL)
        timestamp = datetime.datetime.now().strftime("[%H:%M:%S] ")
        self.log_panel.insert(tk.END, timestamp + message + "\n")
        self.log_panel.config(state=tk.DISABLED)
        self.log_panel.see(tk.END)

    def update_status_bar(self, event=None):
        text_widget = self.get_text_widget()
        if text_widget:
            content = text_widget.get("1.0", tk.END)
            line_count = int(text_widget.index("end-1c").split('.')[0])
            word_count = len(content.split())
            self.status_left.config(text=f"Lines: {line_count} | Words: {word_count}")

    def new_tab(self):
        frame = ttk.Frame(self.tab_control)
        text_widget = ScrolledText(frame, wrap=tk.WORD, font=("Helvetica", 12), undo=True)
        text_widget.pack(fill=tk.BOTH, expand=True)

        text_widget.bind("<<Modified>>", self.on_text_change)

        self.tab_control.add(frame, text="Untitled")
        self.tab_control.select(frame)
        self.tabs[frame] = {"text_widget": text_widget, "file": None}

        self.apply_theme(text_widget)
        self.update_status_bar()
        self.log("New tab created.")

    def on_text_change(self, event):
        text_widget = event.widget
        text_widget.edit_modified(False)
        self.update_status_bar()

    def open_file(self):
        filepath = askopenfilename(filetypes=[("Text Files", "*.txt"), ("All Files", "*.*")])
        if filepath:
            with open(filepath, 'r') as file:
                content = file.read()
            self.new_tab()
            tab = self.get_current_tab()
            tab['text_widget'].insert(tk.END, content)
            tab['file'] = filepath
            tab_name = os.path.basename(filepath)
            self.tab_control.tab(self.tab_control.select(), text=tab_name)
            self.update_status_bar()
            self.log(f"Opened file: {tab_name}")

    def save_as_txt(self):
        tab = self.get_current_tab()
        if not tab:
            return
        filepath = asksaveasfilename(defaultextension=".txt", filetypes=[("Text Files", "*.txt")])
        if filepath:
            with open(filepath, 'w') as f:
                f.write(tab['text_widget'].get(1.0, tk.END).strip())
            tab['file'] = filepath
            self.tab_control.tab(self.tab_control.select(), text=os.path.basename(filepath))
            self.log(f"Saved as TXT: {filepath}")

    def save_as_docx(self):
        tab = self.get_current_tab()
        if not tab:
            return
        filepath = asksaveasfilename(defaultextension=".docx", filetypes=[("Word Documents", "*.docx")])
        if filepath:
            doc = Document()
            doc.add_paragraph(tab['text_widget'].get(1.0, tk.END).strip())
            doc.save(filepath)
            self.log(f"Saved as DOCX: {filepath}")

    def save_as_pdf(self):
        tab = self.get_current_tab()
        if not tab:
            return
        filepath = asksaveasfilename(defaultextension=".pdf", filetypes=[("PDF Files", "*.pdf")])
        if filepath:
            pdf = FPDF()
            pdf.add_page()
            pdf.set_font("Arial", size=12)
            lines = tab['text_widget'].get(1.0, tk.END).strip().split('\n')
            for line in lines:
                pdf.cell(200, 10, txt=line, ln=True)
            pdf.output(filepath)
            self.log(f"Saved as PDF: {filepath}")

    def undo(self):
        try:
            self.get_text_widget().edit_undo()
            self.log("Undo performed.")
        except:
            self.log("Nothing to undo.")

    def redo(self):
        try:
            self.get_text_widget().edit_redo()
            self.log("Redo performed.")
        except:
            self.log("Nothing to redo.")

    def cut(self):
        try:
            self.get_text_widget().event_generate("<<Cut>>")
            self.log("Cut action.")
        except:
            pass

    def copy(self):
        try:
            self.get_text_widget().event_generate("<<Copy>>")
            self.log("Copy action.")
        except:
            pass

    def paste(self):
        try:
            self.get_text_widget().event_generate("<<Paste>>")
            self.log("Paste action.")
        except:
            pass

    def make_bold(self):
        try:
            text = self.get_text_widget()
            text.tag_configure("bold", font=("Helvetica", 12, "bold"))
            if "bold" in text.tag_names(tk.SEL_FIRST):
                text.tag_remove("bold", tk.SEL_FIRST, tk.SEL_LAST)
            else:
                text.tag_add("bold", tk.SEL_FIRST, tk.SEL_LAST)
            self.log("Bold toggled.")
        except:
            self.log("No text selected for bold.")

    def make_italic(self):
        try:
            text = self.get_text_widget()
            text.tag_configure("italic", font=("Helvetica", 12, "italic"))
            if "italic" in text.tag_names(tk.SEL_FIRST):
                text.tag_remove("italic", tk.SEL_FIRST, tk.SEL_LAST)
            else:
                text.tag_add("italic", tk.SEL_FIRST, tk.SEL_LAST)
            self.log("Italic toggled.")
        except:
            self.log("No text selected for italic.")

    def change_text_color(self):
        color = colorchooser.askcolor()[1]
        if color:
            try:
                text = self.get_text_widget()
                tag_name = f"color_{color}"
                text.tag_configure(tag_name, foreground=color)
                text.tag_add(tag_name, tk.SEL_FIRST, tk.SEL_LAST)
                self.log(f"Text color changed to {color}")
            except:
                self.log("No text selected to color.")

    def toggle_theme(self):
        self.dark_mode = not self.dark_mode
        for tab in self.tabs.values():
            self.apply_theme(tab['text_widget'])
        self.log("Theme toggled.")

    def apply_theme(self, text_widget):
        if self.dark_mode:
            text_widget.config(bg="#1e1e1e", fg="#ffffff", insertbackground="white")
        else:
            text_widget.config(bg="white", fg="black", insertbackground="black")

if __name__ == "__main__":
    root = tk.Tk()
    app = AdvancedNotepadApp(root)
    root.mainloop()
